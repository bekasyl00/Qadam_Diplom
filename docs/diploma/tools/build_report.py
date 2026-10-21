"""Rebuild the editable diagrams, PNG exports and report from repository sources."""
from pathlib import Path
import html
import math
import re
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
DIAGRAMS = ROOT / 'diagrams'
DIAGRAMS.mkdir(exist_ok=True)
FONT = Path('/System/Library/Fonts/Supplemental/Arial.ttf')
if not FONT.exists():
    FONT = Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')


class Drawing:
    """Keep editable SVG and raster export in sync without external renderers."""
    def __init__(self, width, height):
        self.w, self.h = width, height
        self.im = Image.new('RGB', (width * 2, height * 2), 'white')
        self.d = ImageDraw.Draw(self.im)
        self.svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">', '<rect width="100%" height="100%" fill="white"/>']

    def text(self, x, y, text, size=22, center=True, bold=False):
        fp = FONT.with_name('Arial Bold.ttf') if bold and FONT.name == 'Arial.ttf' else FONT
        f = ImageFont.truetype(str(fp), size * 2)
        for i, line in enumerate(text.split('\n')):
            yy = y + i * (size + 7)
            box = self.d.textbbox((0, 0), line, font=f)
            xx = x * 2 - (box[2] - box[0]) / 2 if center else x * 2
            self.d.text((xx, yy * 2 - box[1]), line, font=f, fill='#202124')
            self.svg.append(f'<text x="{x}" y="{yy + size}" text-anchor="{"middle" if center else "start"}" font-family="Arial,sans-serif" font-size="{size}" font-weight="{"bold" if bold else "normal"}" fill="#202124">{html.escape(line)}</text>')

    def box(self, x, y, w, h, text='', size=22, fill='#f4f5f7', dash=False):
        self.d.rectangle((x*2,y*2,(x+w)*2,(y+h)*2), fill=fill, outline=None if dash else '#67717e', width=3)
        if dash:
            for a,b in [((x,y),(x+w,y)),((x+w,y),(x+w,y+h)),((x+w,y+h),(x,y+h)),((x,y+h),(x,y))]:
                dx,dy=b[0]-a[0],b[1]-a[1]; length=math.hypot(dx,dy)
                for t in range(0,int(length),13):
                    e=min(t+8,length)
                    self.d.line([(2*(a[0]+dx*t/length),2*(a[1]+dy*t/length)),(2*(a[0]+dx*e/length),2*(a[1]+dy*e/length))],fill='#67717e',width=3)
        dash_attr = 'stroke-dasharray="8 5"' if dash else ''
        self.svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="{fill}" stroke="#67717e" stroke-width="1.5" {dash_attr}/>')
        if text:
            lines = len(text.split('\n'))
            self.text(x+w/2, y+(h-lines*(size+7)+7)/2, text, size)

    def line(self, points, arrow=False, dash=False):
        if not dash:
            self.d.line([(x*2,y*2) for x,y in points], fill='#596372', width=3)
        else:
            for a,b in zip(points,points[1:]):
                dx,dy=b[0]-a[0],b[1]-a[1]; length=math.hypot(dx,dy)
                if not length: continue
                for t in range(0,int(length),13):
                    e=min(t+8,length)
                    self.d.line([(2*(a[0]+dx*t/length),2*(a[1]+dy*t/length)),(2*(a[0]+dx*e/length),2*(a[1]+dy*e/length))],fill='#596372',width=3)
        self.svg.append('<polyline points="'+' '.join(f'{x},{y}' for x,y in points)+'" fill="none" stroke="#596372" stroke-width="1.5"'+(' stroke-dasharray="8 5"' if dash else '')+'/>')
        if arrow:
            x,y = points[-1]; a,b = points[-2]
            ang = math.atan2(y-b,x-a)
            end = [(x,y),(x-13*math.cos(ang-.43),y-13*math.sin(ang-.43)),(x-13*math.cos(ang+.43),y-13*math.sin(ang+.43))]
            self.d.polygon([(a*2,b*2) for a,b in end], fill='#596372')
            self.svg.append('<polygon points="'+' '.join(f'{a},{b}' for a,b in end)+'" fill="#596372"/>')

    def ellipse(self,x,y,w,h,text):
        self.d.ellipse((x*2,y*2,(x+w)*2,(y+h)*2),fill='white',outline='#67717e',width=3)
        self.svg.append(f'<ellipse cx="{x+w/2}" cy="{y+h/2}" rx="{w/2}" ry="{h/2}" fill="white" stroke="#67717e" stroke-width="1.5"/>')
        lines=len(text.split('\n'))
        self.text(x+w/2,y+(h-lines*27+7)/2,text,20)

    def actor(self,x,y,label):
        self.d.ellipse(((x-12)*2,y*2,(x+12)*2,(y+24)*2),outline='#596372',width=3)
        self.svg.append(f'<circle cx="{x}" cy="{y+12}" r="12" fill="white" stroke="#596372" stroke-width="1.5"/>')
        self.line([(x,y+24),(x,y+66)])
        self.line([(x-26,y+40),(x+26,y+40)])
        self.line([(x-22,y+92),(x,y+66),(x+22,y+92)])
        self.text(x,y+106,label,21)

    def save(self,name):
        (DIAGRAMS/f'{name}.svg').write_text('\n'.join(self.svg+['</svg>']))
        self.im.save(DIAGRAMS/f'{name}.png')


def diagrams():
    c=Drawing(1100,590)
    c.box(355,185,390,210,'Qadam platform\nTask formulation and\nteam collaboration',26)
    c.box(25,35,285,110,'Business representative',21)
    c.box(25,430,285,110,'Student team',23)
    c.box(795,35,275,110,'Guest',23)
    c.box(790,430,290,110,'External AI provider\nOpenAI or Groq',23)
    c.line([(310,90),(550,90),(550,185)],True)
    c.text(500,45,'Description, answers, approval',18)
    c.line([(310,480),(440,480),(440,395)],True)
    c.text(520,498,'Proposals and milestone result',18)
    c.line([(870,145),(870,245),(745,245)],True)
    c.text(980,185,'Browse public\nbriefs',17)
    c.line([(745,330),(990,330),(990,430)],True)
    c.text(908,292,'Authorized text / audio',17)
    c.line([(790,475),(710,475),(710,395)],True)
    c.text(831,554,'Questions, suggestions, transcript',17)
    c.save('context')

    c=Drawing(1100,770)
    c.box(260,20,580,725,fill='white')
    c.text(550,37,'Qadam use cases',25,bold=True)
    c.actor(115,155,'Business\nrepresentative')
    c.actor(985,340,'Student team')
    c.actor(115,565,'Guest')
    # Associations drawn behind use cases; labels identify responsibilities.
    for yy in [136,254,372,490]: c.line([(144,205),(300,yy)])
    for yy in [136,372,490,630]: c.line([(956,390),(800,yy)])
    c.line([(144,610),(300,630)])
    c.ellipse(300,90,500,92,'Consult task assistant\nusing accessible context')
    c.ellipse(300,208,500,92,'Formulate, confirm and publish\nwith optional AI and voice input')
    c.ellipse(300,326,500,92,'Submit and review proposal\nTeam submits · business selects')
    c.ellipse(300,444,500,92,'Submit and confirm milestone\nTeam submits · business confirms')
    c.ellipse(300,562,500,136,'Browse published briefs')
    c.save('use-cases')

    c=Drawing(1100,770)
    c.box(15,20,280,685,fill='#fafafa')
    c.text(155,40,'Browser application',23,bold=True)
    c.box(40,100,230,125,'«component»\nVue pages and\nUI components',21)
    c.box(40,280,230,105,'«component»\nPinia workspace',21)
    c.box(40,440,230,105,'«component»\nREST / demo adapter',19)
    c.line([(155,225),(155,280)],True)
    c.line([(155,385),(155,440)],True)
    c.text(155,596,'MediaRecorder captures\noptional audio input',18)
    c.box(350,20,440,685,fill='#fafafa')
    c.text(570,40,'Go application',23,bold=True)
    c.box(385,100,370,95,'«component» HTTP API\nValidation and demo-actor access',19)
    c.box(385,250,370,100,'«component» Domain services\nTasks · proposals · milestones · teams',18)
    c.box(385,405,175,110,'«component»\nRating\nPure calculation',18)
    c.box(585,405,170,110,'«component»\nAI adapter\nSource checks',18)
    c.box(385,575,370,80,'«component» SQL repositories',19)
    c.line([(270,487),(322,487),(322,147),(385,147)],True)
    c.text(320,212,'REST',16)
    c.line([(570,195),(570,250)],True)
    c.line([(475,350),(475,405)],True)
    c.line([(735,195),(775,195),(775,375),(670,375),(670,405)],True)
    c.line([(570,350),(570,575)],True)
    c.box(840,565,245,100,'PostgreSQL\nTasks, snapshots,\nproposals and awards',19)
    c.box(840,380,245,140,'External AI\nOpenAI / Groq\nText and transcription',20)
    c.line([(755,615),(840,615)],True)
    c.line([(755,460),(840,460)],True)
    c.text(550,730,'Client-selected demo identity is not production authentication',20)
    c.save('components')

    c=Drawing(1100,475)
    labels=[('01 Describe','Business description\n[ Editable text ]\n[ Record voice ]\nGet AI questions'),('02 Clarify','Contextual questions\n[ Editable answers ]\nSource-linked input\nCreate draft brief'),('03 Review','16 editable fields\nCompleteness indicator\nReview suggestions\nConfirm and publish'),('04 Collaborate','Published task\nTeam proposal\nBusiness selection\nConfirm milestone')]
    for i,(title,body) in enumerate(labels):
        x=20+i*275
        c.box(x,25,240,315,fill='white')
        c.box(x,25,240,64,title,23)
        for j,line in enumerate(body.split('\n')):
            c.text(x+120,125+j*49,line,19)
        if i<3:c.line([(x+240,195),(x+272,195)],True)
    c.text(550,380,'AI unavailable → manual form remains available',23)
    c.text(550,430,'Schematic workflow, not an application screenshot',20)
    c.save('workflow')


def plain(s):
    return re.sub(r'\*\*(.*?)\*\*',r'\1',s).replace('`','')


def hyperlink(p, label, target):
    from docx.opc.constants import RELATIONSHIP_TYPE
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), p.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1'); props.append(color)
    underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
    run.append(props)
    text = OxmlElement('w:t'); text.text = label; run.append(text)
    link.append(run); p._p.append(link)


def runs(p,s):
    pattern = r'(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|https?://[^\s<>]+)'
    for token in re.split(pattern,s):
        match = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', token)
        if match:
            label, target = match.groups()
            if not target.startswith(('https://', 'http://')):
                target = 'https://github.com/bekasyl00/Qadam_Diplom/blob/QADAM-12-predefense-revision/docs/diploma/' + target
            hyperlink(p, label, target)
        elif token.startswith(('https://', 'http://')):
            target = token.rstrip('.,;')
            hyperlink(p, target, target)
            if len(target) < len(token): p.add_run(token[len(target):])
        else:
            r=p.add_run(plain(token))
            if token.startswith('**'):r.bold=True
            if token.startswith('`'):r.font.name='Arial'; r.font.size=Pt(10)


def report():
    doc=Document()
    for border in list(doc.styles.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
    sec=doc.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.65)
    sec.left_margin=sec.right_margin=Inches(.7)
    sec.header_distance=sec.footer_distance=Inches(.3)
    sec.different_first_page_header_footer=True
    for name in ['Normal','Body Text','List Bullet','List Number']:
        style=doc.styles[name]
        style.font.name='Arial'; style.font.size=Pt(11)
        style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_after=Pt(7)
        style.paragraph_format.line_spacing=1.12
        style.paragraph_format.widow_control=True
    for name,size in [('Title',27),('Heading 1',18),('Heading 2',12)]:
        style=doc.styles[name]; style.font.name='Arial'; style.font.size=Pt(size)
        style.font.bold=True; style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_before=Pt(12 if name=='Heading 2' else 0)
        style.paragraph_format.space_after=Pt(10)
        style.paragraph_format.keep_with_next=True
    h=sec.header.paragraphs[0]; h.text='Qadam   |   Diploma pre-defense preparation'
    h.style=doc.styles['Normal']
    for r in h.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0)
    f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    f.add_run('Qadam  ·  ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
    for r in f.runs:r.font.size=Pt(9)
    lines=(ROOT/'REPORT.md').read_text().splitlines()
    i=0
    page_start=False
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1;continue
        if line=='<!-- page -->':page_start=True;i+=1;continue
        if line.startswith('# '):
            p=doc.add_paragraph(line[2:],style='Title');p.paragraph_format.space_before=Pt(50);p.paragraph_format.space_after=Pt(26)
        elif line.startswith('## '):
            p=doc.add_paragraph(line[3:],style='Heading 1')
            p.paragraph_format.page_break_before=page_start
            page_start=False
        elif line.startswith('### '):doc.add_paragraph(line[4:],style='Heading 2')
        elif line.startswith('!['):
            match=re.match(r'!\[(.*?)\]\((.*?)\)',line)
            doc.add_picture(str(ROOT/match.group(2)),width=Inches(6.35 if any(n in line for n in ['components', 'activity', 'sequence']) else 7.05))
            doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
            p=doc.add_paragraph(match.group(1));p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:r.font.size=Pt(9);r.italic=True
        elif line.startswith('|'):
            data=[]
            while i<len(lines) and lines[i].startswith('|'):
                row=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[- :]+',x) for x in row):data.append(row)
                i+=1
            n=len(data[0]); table=doc.add_table(rows=0,cols=n)
            table.autofit=False
            # Unequal columns allocate more room to analytical text.
            widths={2:[1.25,5.85],3:[1.48,2.75,2.87],4:[1.30,1.92,1.94,1.94],5:[1.15,1.48,1.49,1.49,1.49]}[n]
            if n==3 and len(data)>7:widths=[.97,2.62,3.51]
            widths={
                'Hypothesis':[2.0,3.25,1.85],
                'Team member':[1.25,2.9,2.95],
                'Supervisor recommendation':[1.65,3.0,2.45],
                'Finding or limitation':[2.25,2.45,2.4],
                'Publication and reading basis':[1.6,3.0,2.5],
                'Decision':[1.3,3.2,2.6],
                'User goal':[1.3,1.8,4.0],
                'Member':[2.05,3.7,1.35],
                'Risk':[1.5,4.6,1.0],
                'Assignment component':[1.4,2.7,3.0],
            }.get(data[0][0],widths)
            for col,w in zip(table.columns,widths):col.width=Inches(w)
            props=table._tbl.tblPr
            borders=OxmlElement('w:tblBorders')
            for side in ['top','left','bottom','right','insideH','insideV']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D5D8DE');borders.append(el)
            props.append(borders)
            margins=OxmlElement('w:tblCellMar')
            for side,value in [('top','85'),('bottom','85'),('left','90'),('right','90')]:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),value);el.set(qn('w:type'),'dxa');margins.append(el)
            props.append(margins)
            for ri,values in enumerate(data):
                row=table.add_row()
                trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
                if ri==0:trpr.append(OxmlElement('w:tblHeader'))
                for j,(cell,text) in enumerate(zip(row.cells,values)):
                    cell.width=Inches(widths[j])
                    p=cell.paragraphs[0];runs(p,text)
                    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.04
                    for r in p.runs:r.font.size=Pt(9.5 if n>=5 else 10);r.bold=(ri==0 or r.bold)
                    if ri==0:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F0F1F3');cell._tc.get_or_add_tcPr().append(shade)
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1;p.add_run().font.size=Pt(2)
            continue
        elif line.startswith('- '):runs(doc.add_paragraph(style='List Bullet'),line[2:])
        elif re.match(r'^\d+\. ',line):
            p=doc.add_paragraph()
            p.paragraph_format.left_indent=Pt(18)
            p.paragraph_format.first_line_indent=Pt(-18)
            runs(p,line)
        else:runs(doc.add_paragraph(),line)
        i+=1
    doc.core_properties.title='Development of the Qadam Platform for Business Task Formulation Using Artificial Intelligence'
    doc.core_properties.author='Bekasyl Zhaksylyk; Mereke Myrzakhan; Amantai Zere'
    doc.core_properties.subject='Pre-defense preparation for 24 October; supervisor Altynbek Seitenov'
    doc.save(ROOT/'Qadam_Predefense_Report.docx')


if __name__=='__main__':
    diagrams()
    from architecture_diagrams import create_architecture_diagrams
    create_architecture_diagrams()
    report()
    print('Generated diagrams and Qadam_Predefense_Report.docx')
